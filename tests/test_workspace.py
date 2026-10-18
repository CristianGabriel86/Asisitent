from io import BytesIO
import re

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from aistent.db import db


def upload(client, project_id, name='notite.txt', data=b'Text pentru analiza'):
    response = client.post(f'/projects/{project_id}/upload', data={'files': (BytesIO(data), name)}, content_type='multipart/form-data')
    assert response.status_code == 302
    with client.application.app_context():
        return db().execute('SELECT * FROM files WHERE project_id=? ORDER BY rowid DESC LIMIT 1', (project_id,)).fetchone()['id']


def test_project_and_multi_upload_and_edit(client, project_id):
    response = client.get(f'/projects/{project_id}')
    assert response.status_code == 200
    assert b'Biblioteca proiectului' in response.data
    response = client.post(f'/projects/{project_id}/upload', data={'files': [(BytesIO(b'Primul'), 'primul.txt'), (BytesIO(b'Al doilea'), 'al-doilea.txt')]})
    assert response.status_code == 302
    with client.application.app_context():
        files = db().execute('SELECT * FROM files').fetchall()
        assert len(files) == 2
    assert client.get(f"/files/{files[0]['id']}/edit").status_code == 200
    assert client.get(f"/files/{files[0]['id']}").data == b'Primul'


def test_files_and_projects_are_private(app, client, project_id):
    file_id = upload(client, project_id)
    other = app.test_client()
    other.post('/register', data={'email': 'other@example.com', 'password': 'another-long-password'})
    other.post('/login', data={'email': 'other@example.com', 'password': 'another-long-password'})
    for path in (f'/projects/{project_id}', f'/files/{file_id}', f'/files/{file_id}/edit'):
        assert other.get(path).status_code == 404
    assert other.post(f'/files/{file_id}/delete').status_code == 404
    assert other.post(f'/projects/{project_id}/documents', data={'title': 'Test', 'content': 'Text'}).status_code == 404
    anonymous = app.test_client()
    assert anonymous.get(f'/files/{file_id}').status_code == 302
    assert anonymous.post('/documents', data={'title': 'Test', 'content': 'Text'}).status_code == 302


def test_csrf_required_and_password_hashed(app):
    app.config['CSRF_DISABLED_FOR_TESTS'] = False
    client = app.test_client()
    page = client.get('/register')
    assert page.status_code == 200
    assert client.post('/register', data={'email': 'csrf@example.com', 'password': 'long-password-here'}).status_code == 400
    token = re.search(rb'name="csrf_token" value="([^"]+)"', page.data).group(1).decode()
    assert client.post('/register', data={'email': 'csrf@example.com', 'password': 'long-password-here', 'csrf_token': token}).status_code == 302
    with app.app_context():
        password = db().execute('SELECT password FROM users').fetchone()[0]
        assert password.startswith('scrypt:') and password != 'long-password-here'


def test_upload_rejects_fake_image_and_scripts(client, project_id):
    for name in ('attack.html', 'fake.png'):
        assert client.post(f'/projects/{project_id}/upload', data={'files': (BytesIO(b'<script>'), name)}).status_code == 400


def test_docx_import_edit_export_round_trip(client, project_id):
    output = BytesIO()
    doc = Document()
    doc.add_heading('Introducere', 1)
    doc.add_paragraph('Conținut editabil.')
    doc.save(output)
    file_id = upload(client, project_id, 'lucrare.docx', output.getvalue())
    response = client.get(f'/files/{file_id}/edit')
    assert response.status_code == 200
    assert '## Introducere' in response.text and 'Conținut editabil.' in response.text
    response = client.post(f'/projects/{project_id}/documents', data={'title': 'Revizuit', 'content': 'Text modificat', 'format': 'docx'})
    assert response.status_code == 302
    with client.application.app_context():
        result_id = db().execute("SELECT id FROM files WHERE name='document.docx'").fetchone()[0]
    document = Document(BytesIO(client.get(f'/files/{result_id}').data))
    assert document.paragraphs[1].text == 'Text modificat'


def test_presentation_export_and_import(client, project_id):
    response = client.post('/documents', data={'title': 'Prezentare', 'content': '## Obiective\nPrimul obiectiv\n## Concluzii\nRezultate', 'format': 'pptx'})
    assert response.status_code == 200
    slides = Presentation(BytesIO(response.data)).slides
    assert len(slides) == 3
    assert slides[1].shapes.title.text == 'Obiective'
    file_id = upload(client, project_id, 'prezentare.pptx', response.data)
    assert 'Primul obiectiv' in client.get(f'/files/{file_id}/edit').text


def test_electrical_estimate_totals(client, project_id):
    response = client.post(f'/projects/{project_id}/documents', data={'title': 'Deviz', 'content': 'Descriere,UM,Cantitate,Material,Manopera\nCablu,m,10,4.5,3\nPriza,buc,2,25,30', 'format': 'estimate', 'vat': '21'})
    assert response.status_code == 302
    with client.application.app_context():
        file_id = db().execute('SELECT id FROM files').fetchone()[0]
    sheet = load_workbook(BytesIO(client.get(f'/files/{file_id}').data)).active
    assert sheet['H2'].value == 75
    assert sheet['H3'].value == 110
    assert sheet['H4'].value == 185
    assert sheet['H5'].value == 38.85
    assert sheet['H6'].value == 223.85


def test_photo_edit_and_private_preview(client, project_id):
    image = Image.new('RGBA', (24, 24), (0, 0, 0, 0))
    output = BytesIO()
    image.save(output, 'PNG')
    file_id = upload(client, project_id, 'test.png', output.getvalue())
    response = client.post(f'/projects/{project_id}/photo', data={'file_ids': file_id, 'brightness': '1', 'saturation': '1', 'background': '#ff0000'})
    assert response.status_code == 302
    with client.application.app_context():
        result = db().execute("SELECT id FROM files WHERE name='fotografie-editata.png'").fetchone()[0]
    response = client.get(f'/files/{result}?preview=1')
    assert 'inline' in response.headers['Content-Disposition']
    assert Image.open(BytesIO(response.data)).getpixel((0, 0)) == (255, 0, 0, 255)


def test_delete_own_file(client, project_id):
    file_id = upload(client, project_id)
    assert client.post(f'/files/{file_id}/delete').status_code == 302
    assert client.get(f'/files/{file_id}').status_code == 404


def test_numeric_analysis(client):
    response = client.post('/documents', data={'title': 'Analiză', 'content': 'Lună,Venit\nIanuarie,1200\nFebruarie,1500', 'format': 'analysis'})
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.data)).active
    assert sheet['B2'].value == 2
    assert sheet['B3'].value == 2700
    assert sheet['B4'].value == 1350
    assert sheet['B5'].value == 1200
    assert sheet['B6'].value == 1500


def test_invalid_estimate_and_analysis(client, project_id):
    for kind, content in [('estimate', 'A,B,C,D,E\nCablu,m,-1,4,2'), ('analysis', 'Eticheta,Valoare\nA,nan')]:
        response = client.post(f'/projects/{project_id}/documents', data={'title': 'Test', 'content': content, 'format': kind, 'vat': '21'})
        assert response.status_code == 400


def test_login_throttle_and_closed_registration(app):
    visitor = app.test_client()
    for _ in range(10):
        assert visitor.post('/login', data={'email': 'missing@example.com', 'password': 'wrong'}).status_code == 400
    assert visitor.post('/login', data={'email': 'missing@example.com', 'password': 'wrong'}).status_code == 429
    app.config['ALLOW_REGISTRATION'] = False
    assert visitor.get('/register').status_code == 403
