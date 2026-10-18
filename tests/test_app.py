from io import BytesIO

from docx import Document
from aistent import create_app


def test_home():
    response = create_app().test_client().get('/')
    assert response.status_code == 200
    assert b'Aistent' in response.data


def test_download_word_preserves_romanian_and_paragraphs():
    response = create_app().test_client().post('/documents', data={
        'title': 'Cerere de concediu', 'content': 'Bună ziua!\nVă mulțumesc.'
    })
    assert response.status_code == 200
    assert 'attachment' in response.headers['Content-Disposition']
    document = Document(BytesIO(response.data))
    assert [p.text for p in document.paragraphs] == ['Cerere de concediu', 'Bună ziua!', 'Vă mulțumesc.']


def test_invalid_input():
    client = create_app().test_client()
    for data in ({}, {'title': 'Titlu', 'content': ' '}, {'title': 'x' * 201, 'content': 'Text'}, {'title': 'Titlu', 'content': '\x00'}):
        assert client.post('/documents', data=data).status_code == 400


def test_request_size_limit():
    response = create_app().test_client().post('/documents', data={'title': 'Titlu', 'content': 'x' * 300000})
    assert response.status_code == 413


def test_pdf_text():
    from pypdf import PdfReader
    response = create_app().test_client().post('/documents', data={'title': 'Știință', 'content': 'Cercetare și educație <test>', 'format': 'pdf'})
    assert response.status_code == 200
    text = ''.join(page.extract_text() for page in PdfReader(BytesIO(response.data)).pages)
    assert 'Știință' in text
    assert 'Cercetare și educație <test>' in text


def test_excel_literal_cells():
    from openpyxl import load_workbook
    response = create_app().test_client().post('/documents', data={'title': 'Tabel', 'content': 'Nume,Valoare\nAna,12\nFormula,=1+1', 'format': 'xlsx'})
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.data)).active
    assert sheet['A2'].value == 'Ana'
    assert sheet['B3'].value == '=1+1'
    assert sheet['B3'].data_type == 's'


def test_chart_is_readable_image():
    from PIL import Image
    response = create_app().test_client().post('/documents', data={'title': 'Venit', 'content': 'Lună,Lei\nIanuarie,1200\nFebruarie,1500', 'format': 'chart'})
    assert response.status_code == 200
    image = Image.open(BytesIO(response.data))
    image.verify()
    assert image.format == 'PNG'


def test_academic_structure_and_explicit_headings():
    client = create_app().test_client()
    response = client.post('/documents', data={'title': 'Studiu', 'content': 'Text propriu.', 'format': 'academic'})
    paragraphs = [p.text for p in Document(BytesIO(response.data)).paragraphs]
    assert 'Metodologie' in paragraphs and 'Bibliografie' in paragraphs
    assert 'Text propriu.' in paragraphs
    response = client.post('/documents', data={'title': 'Studiu', 'content': '## Rezumat\nUn rezumat.', 'format': 'academic'})
    document = Document(BytesIO(response.data))
    assert document.paragraphs[1].style.name == 'Heading 1'
    assert document.paragraphs[2].text == 'Un rezumat.'


def test_invalid_exports():
    client = create_app().test_client()
    for kind, content in [('unknown', 'Text'), ('xlsx', 'A,B\nC'), ('chart', 'A,B\nC,abc'), ('chart', 'A,B\nC,nan'), ('chart', 'A,B,C\n1,2,3')]:
        assert client.post('/documents', data={'title': 'Test', 'content': content, 'format': kind}).status_code == 400
