"""Aplicația Aistent pentru generarea documentelor Word."""
from io import BytesIO

from docx import Document
from flask import Flask, render_template, request, send_file


def create_app():
    app = Flask(__name__)
    app.config['MAX_CONTENT_LENGTH'] = 256 * 1024

    @app.get('/')
    def index():
        return render_template('index.html')

    @app.post('/documents')
    def generate_document():
        title = request.form.get('title', '').strip()
        content = request.form.get('content', '').strip()
        if not title or not content:
            return render_template('index.html', error='Completează titlul și conținutul.', title=title, content=content), 400
        if len(title) > 200 or len(content) > 100_000:
            return render_template('index.html', error='Titlul poate avea cel mult 200 de caractere, iar conținutul 100.000.', title=title, content=content), 400
        # XML 1.0 does not support these control characters.
        if any(ord(char) < 32 and char not in '\n\r\t' for char in title + content):
            return render_template('index.html', error='Textul conține caractere de control nepermise.', title=title, content=content), 400
        document = Document()
        document.add_heading(title, level=0)
        for paragraph in content.splitlines():
            document.add_paragraph(paragraph)
        output = BytesIO()
        document.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name='document.docx', mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

    @app.errorhandler(413)
    def too_large(error):
        return render_template('index.html', error='Documentul trimis este prea mare.'), 413

    return app
