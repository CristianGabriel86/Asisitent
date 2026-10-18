"""Exporturi locale: nu necesită servicii sau chei API."""
import csv
import math
import os
import tempfile
from pathlib import Path
from io import BytesIO, StringIO
from threading import Lock
from xml.sax.saxutils import escape

from docx import Document
os.environ.setdefault('MPLCONFIGDIR', str(Path(tempfile.gettempdir()) / f'aistent-matplotlib-{os.getuid() if hasattr(os, "getuid") else "user"}'))

from matplotlib import get_data_path
from matplotlib.figure import Figure
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

_PLOT_LOCK = Lock()
pdfmetrics.registerFont(TTFont('DejaVuSans', str(Path(get_data_path()) / 'fonts/ttf/DejaVuSans.ttf')))
SECTIONS = ('Rezumat', 'Introducere', 'Revizuirea literaturii', 'Metodologie', 'Rezultate', 'Discuții', 'Concluzii', 'Bibliografie')


def word(title, content, academic=False):
    document = Document()
    document.add_heading(title, 0)
    if academic and not any(line.startswith('## ') for line in content.splitlines()):
        for section in SECTIONS:
            document.add_heading(section, 1)
            document.add_paragraph(content if section == 'Introducere' else '[De completat de autor]')
    else:
        for line in content.splitlines():
            if academic and line.startswith('## '):
                document.add_heading(line[3:], 1)
            else:
                document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def pdf(title, content):
    styles = getSampleStyleSheet()
    body = ParagraphStyle('RomanianBody', parent=styles['BodyText'], fontName='DejaVuSans', leading=15, wordWrap='CJK')
    heading = ParagraphStyle('RomanianTitle', parent=body, fontSize=18, leading=23, spaceAfter=18)
    output = BytesIO()
    story = [Paragraph(escape(title), heading)]
    for line in content.splitlines():
        story.extend([Paragraph(escape(line) or '&#160;', body), Spacer(1, 6)])
    SimpleDocTemplate(output).build(story)
    return output.getvalue()


def table(content):
    try:
        rows = list(csv.reader(StringIO(content), strict=True))
    except csv.Error as error:
        raise ValueError('Date CSV invalide.') from error
    if len(rows) < 2 or not rows[0] or len(rows[0]) > 50 or len(rows) > 1000:
        raise ValueError('Introdu un antet și cel puțin un rând; maximum 1.000 de rânduri și 50 de coloane.')
    if any(len(row) != len(rows[0]) for row in rows):
        raise ValueError('Toate rândurile CSV trebuie să aibă același număr de coloane.')
    return rows


def excel(title, content):
    rows = table(content)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Date'
    for row in rows:
        sheet.append(row)
    # Treat all user-provided cells as literal text, including spreadsheet formulas.
    for row in sheet:
        for cell in row:
            cell.data_type = 's'
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def chart(title, content):
    rows = table(content)
    if len(rows[0]) != 2 or len(rows) > 101:
        raise ValueError('Graficul necesită exact două coloane și cel mult 100 de valori.')
    try:
        values = [float(row[1]) for row in rows[1:]]
    except ValueError as error:
        raise ValueError('A doua coloană trebuie să conțină numere.') from error
    if not all(math.isfinite(value) for value in values):
        raise ValueError('Valorile trebuie să fie numere finite.')
    with _PLOT_LOCK:
        figure = Figure(figsize=(10, 6), layout='constrained')
        axis = figure.subplots()
        axis.bar(range(len(values)), values, tick_label=[row[0] for row in rows[1:]])
        axis.set(title=title, xlabel=rows[0][0], ylabel=rows[0][1])
        axis.tick_params(axis='x', labelrotation=45)
        output = BytesIO()
        figure.savefig(output, format='png', dpi=150)
    return output.getvalue()


EXPORTS = {
    'docx': (word, 'document.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
    'pdf': (pdf, 'document.pdf', 'application/pdf'),
    'xlsx': (excel, 'tabel.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
    'chart': (chart, 'grafic.png', 'image/png'),
    'academic': (lambda title, content: word(title, content, academic=True), 'lucrare-academica.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
}
