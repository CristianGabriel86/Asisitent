"""Import editabil, prezentări și devize electrice calculate cu Decimal."""
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pypdf import PdfReader

from .exports import EXPORTS, table


def presentation(title, content):
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[0])
    first.shapes.title.text = title
    first.placeholders[1].text = 'Creat cu Aistent'
    sections = content.split('\n## ')
    if len(sections) > 40:
        raise ValueError('Prezentarea poate avea maximum 40 de secțiuni.')
    for section in sections:
        lines = section.removeprefix('## ').splitlines()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = lines[0][:200] if lines else 'Secțiune'
        slide.placeholders[1].text = '\n'.join(lines[1:]) or section
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def extract(path, name):
    suffix = Path(name).suffix.lower()
    if suffix in {'.txt', '.csv', '.srt'}:
        with Path(path).open(encoding='utf-8-sig') as source:
            return source.read(100_000)
    if suffix == '.docx':
        document = Document(path)
        lines = []
        for paragraph in document.paragraphs:
            prefix = '## ' if paragraph.style and paragraph.style.name.startswith('Heading') else ''
            lines.append(prefix + paragraph.text)
        for grid in document.tables:
            lines.extend(' | '.join(cell.text for cell in row.cells) for row in grid.rows)
        return '\n'.join(lines)[:100_000]
    if suffix == '.pdf':
        reader = PdfReader(path)
        if len(reader.pages) > 200:
            raise ValueError('PDF-ul are mai mult de 200 de pagini.')
        text = ''
        for page in reader.pages:
            text += (page.extract_text() or '') + '\n'
            if len(text) >= 100_000:
                break
        if not text.strip():
            raise ValueError('PDF-ul este scanat; OCR nu este inclus în această versiune.')
        return text[:100_000]
    if suffix == '.xlsx':
        import csv
        from io import StringIO
        workbook = load_workbook(path, read_only=True, data_only=False)
        output = StringIO()
        writer = csv.writer(output)
        sheet = workbook.active
        for row in sheet.iter_rows(max_row=min(sheet.max_row or 0, 1000), max_col=min(sheet.max_column or 0, 50), values_only=True):
            writer.writerow(['' if cell is None else cell for cell in row])
        workbook.close()
        return output.getvalue()[:100_000]
    if suffix == '.pptx':
        return '\n'.join('## ' + slide.shapes.title.text + '\n' + '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame and shape != slide.shapes.title) if slide.shapes.title else '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in Presentation(path).slides)[:100_000]
    raise ValueError('Acest format nu poate fi importat ca text.')


def estimate(content, vat='19'):
    """CSV: descriere,unitate,cantitate,material_unitar,manopera_unitara."""
    rows = table(content)
    if len(rows[0]) != 5:
        raise ValueError('Devizul necesită 5 coloane: descriere, unitate, cantitate, material unitar, manoperă unitară.')
    try:
        rate = Decimal(vat)
        if not rate.is_finite() or not 0 <= rate <= 100:
            raise ValueError('TVA trebuie să fie între 0 și 100%.')
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Deviz electric'
        sheet.append(['Descriere', 'UM', 'Cantitate', 'Material / UM', 'Manoperă / UM', 'Total material', 'Total manoperă', 'Total net'])
        total = Decimal('0')
        cent = Decimal('0.01')
        for row in rows[1:]:
            quantity, material, labor = [Decimal(value) for value in row[2:]]
            if any(not value.is_finite() or value < 0 or value > Decimal('1000000000') for value in (quantity, material, labor)):
                raise ValueError('Valorile din deviz trebuie să fie finite, pozitive și cel mult un miliard.')
            materials = (quantity * material).quantize(cent, rounding=ROUND_HALF_UP)
            work = (quantity * labor).quantize(cent, rounding=ROUND_HALF_UP)
            subtotal = materials + work
            total += subtotal
            sheet.append([row[0], row[1], float(quantity), float(material), float(labor), float(materials), float(work), float(subtotal)])
            for cell in sheet[sheet.max_row][:2]:
                cell.data_type = 's'
        tax = (total * rate / 100).quantize(cent, rounding=ROUND_HALF_UP)
        sheet.append(['Total net', '', '', '', '', '', '', float(total)])
        sheet.append([f'TVA {rate}%', '', '', '', '', '', '', float(tax)])
        sheet.append(['TOTAL CU TVA', '', '', '', '', '', '', float(total + tax)])
        sheet.freeze_panes = 'A2'
        for row in sheet.iter_rows(min_row=2, min_col=3):
            for cell in row:
                cell.number_format = '#,##0.00'
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    except InvalidOperation as error:
        raise ValueError('Cantitățile, prețurile și TVA trebuie să fie numere valide.') from error


EXPORTS['pptx'] = (presentation, 'prezentare.pptx', 'application/vnd.openxmlformats-officedocument.presentationml.presentation')


def analysis(title, content):
    rows = table(content)
    if len(rows[0]) != 2:
        raise ValueError('Analiza necesită două coloane: etichetă și valoare numerică.')
    try:
        values = [Decimal(row[1]) for row in rows[1:]]
        if any(not value.is_finite() or abs(value) > Decimal('1e15') for value in values):
            raise ValueError('Valorile trebuie să fie finite, cu modulul cel mult 10^15.')
    except InvalidOperation as error:
        raise ValueError('A doua coloană trebuie să conțină numere.') from error
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Analiză'
    sheet.append(['Indicator', 'Valoare'])
    for label, value in [('Număr observații', len(values)), ('Sumă', sum(values)), ('Medie', sum(values) / len(values)), ('Minim', min(values)), ('Maxim', max(values))]:
        sheet.append([label, float(value)])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


EXPORTS['analysis'] = (analysis, 'analiza.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
